"""Private visitor contact intake and Excel export for the public site."""
from __future__ import annotations

import json
import os
import re
import sqlite3
import subprocess
import threading
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import load_json

REPOSITORY = "alexanderh2seo4/afs-contact-submissions"
EXPORT_NAME = "visitor-submissions.xlsx"
FIELDS = {"submissionId", "name", "email", "phone", "postalCode", "city", "interests", "consent", "consentVersion", "website"}
INTEREST_LABELS = {
    "sending_interviews": "Sending-Interviews",
    "hosting": "Hosting und Gastfamilien",
    "exchange": "Austauschjahr",
    "other": "Sonstiges",
}
EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
UUID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[1-8][0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}$", re.I)


class ContactIntakeError(ValueError):
    """A safe, constant error code for invalid or unavailable intake."""


class ContactIntake:
    def __init__(self, directory: Path):
        self.directory = directory
        self.config_path = directory / "contact-intake.json"
        self.database = directory / "visitor-contacts.sqlite3"
        if self.database.is_symlink():
            raise ContactIntakeError("contact_database_unavailable")
        directory.mkdir(mode=0o700, parents=True, exist_ok=True)
        directory.chmod(0o700)
        fd = os.open(self.database, os.O_CREAT | os.O_RDWR, 0o600)
        os.close(fd)
        self.database.chmod(0o600)
        self.lock = threading.RLock()
        with self.connect() as db:
            db.execute("""CREATE TABLE IF NOT EXISTS submissions (
                submission_id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                name TEXT NOT NULL,
                email TEXT NOT NULL,
                phone TEXT NOT NULL,
                postal_code TEXT NOT NULL,
                city TEXT NOT NULL,
                interests TEXT NOT NULL DEFAULT '[]',
                consent_at TEXT NOT NULL,
                consent_version TEXT NOT NULL
            )""")
            columns = {row["name"] for row in db.execute("PRAGMA table_info(submissions)")}
            if "interests" not in columns:
                db.execute("ALTER TABLE submissions ADD COLUMN interests TEXT NOT NULL DEFAULT '[]'")

    @property
    def config(self) -> dict:
        try:
            value = load_json(self.config_path, {})
        except Exception:
            return {}
        return value if isinstance(value, dict) else {}

    @property
    def public_config(self) -> dict:
        path = self.directory.parent / "docs" / "assets" / "contact-config.json"
        try:
            value = json.loads(path.read_text())
        except Exception:
            return {}
        return value if isinstance(value, dict) else {}

    @property
    def enabled(self) -> bool:
        config = self.config
        public = self.public_config
        endpoint = public.get("apiBaseUrl", "")
        try:
            if not isinstance(endpoint, str):
                raise ValueError()
            parsed_endpoint = urlsplit(endpoint)
            parsed_endpoint.port
            valid_endpoint = parsed_endpoint.scheme == "https" and bool(parsed_endpoint.hostname) and not parsed_endpoint.username and not parsed_endpoint.password and parsed_endpoint.path in {"", "/"} and not parsed_endpoint.query and not parsed_endpoint.fragment
        except Exception:
            valid_endpoint = False
        required = ("noticeVersion", "purpose", "retentionText", "privacyContact")
        return bool(
            config.get("enabled") is True
            and config.get("repository") == REPOSITORY
            and all(isinstance(config.get(field), str) and config[field].strip() for field in required)
            and config.get("repositoryPath")
            and public.get("enabled") is True
            and valid_endpoint
            and all(config.get(field) == public.get(field) for field in required)
        )

    @contextmanager
    def connect(self):
        import contextlib

        with contextlib.closing(sqlite3.connect(self.database, timeout=30)) as db:
            db.execute("PRAGMA busy_timeout = 30000")
            db.row_factory = sqlite3.Row
            with db:
                yield db

    @staticmethod
    def _text(value, limit: int, field: str) -> str:
        if not isinstance(value, str):
            raise ContactIntakeError("invalid_submission")
        value = " ".join(value.strip().split())
        if not value or len(value) > limit or any(ord(char) < 32 for char in value):
            raise ContactIntakeError("invalid_submission")
        return value

    def validate(self, payload: object) -> dict | None:
        if not self.enabled:
            raise ContactIntakeError("intake_unavailable")
        if not isinstance(payload, dict) or set(payload) != FIELDS:
            raise ContactIntakeError("invalid_submission")
        # The honeypot is deliberately accepted without storing anything.
        if payload.get("website"):
            return None
        submission_id = payload.get("submissionId")
        if not isinstance(submission_id, str) or not UUID.fullmatch(submission_id):
            raise ContactIntakeError("invalid_submission")
        if payload.get("consent") is not True or payload.get("consentVersion") != self.config.get("noticeVersion"):
            raise ContactIntakeError("invalid_submission")
        name = self._text(payload.get("name"), 120, "name")
        email = self._text(payload.get("email"), 254, "email")
        phone = self._text(payload.get("phone"), 32, "phone")
        postal_code = self._text(payload.get("postalCode"), 5, "postal_code")
        city = self._text(payload.get("city"), 100, "city")
        interests = payload.get("interests")
        if not isinstance(interests, list) or len(interests) > len(INTEREST_LABELS) or any(not isinstance(value, str) or value not in INTEREST_LABELS for value in interests) or len(set(interests)) != len(interests):
            raise ContactIntakeError("invalid_submission")
        interests = [value for value in INTEREST_LABELS if value in interests]
        digits = sum(char.isdigit() for char in phone)
        if not EMAIL.fullmatch(email) or not re.fullmatch(r"\d{5}", postal_code):
            raise ContactIntakeError("invalid_submission")
        if digits < 7 or digits > 15 or not re.fullmatch(r"[+0-9() .-]+", phone):
            raise ContactIntakeError("invalid_submission")
        accepted_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        return {
            "submission_id": submission_id.lower(),
            "created_at": accepted_at,
            "name": name,
            "email": email,
            "phone": phone,
            "postal_code": postal_code,
            "city": city,
            "interests": json.dumps(interests, ensure_ascii=False),
            "consent_at": accepted_at,
            "consent_version": self.config["noticeVersion"],
        }

    def submit(self, payload: object) -> bool:
        row = self.validate(payload)
        if row is None:
            return False
        with self.lock, self.connect() as db:
            db.execute("""INSERT OR IGNORE INTO submissions
                (submission_id, created_at, name, email, phone, postal_code, city, interests, consent_at, consent_version)
                VALUES (:submission_id, :created_at, :name, :email, :phone, :postal_code, :city, :interests, :consent_at, :consent_version)""", row)
        return True

    def records(self) -> list[sqlite3.Row]:
        with self.connect() as db:
            return db.execute("SELECT * FROM submissions ORDER BY created_at, submission_id").fetchall()

    @staticmethod
    def _excel_text(value: str) -> str:
        # Keep visitor-entered text literal when a workbook is opened in Excel.
        return "'" + value if value[:1] in "=+-@" else value

    def export_workbook(self, output: Path) -> None:
        records = self.records()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Besucher-Kontakte"
        sheet.append(["Eingegangen (UTC)", "Name", "E-Mail", "Telefonnummer", "Postleitzahl", "Ort", "Interessen", "Einwilligung (UTC)", "Hinweisversion"])
        for record in records:
            try:
                interests = json.loads(record["interests"])
            except (TypeError, json.JSONDecodeError):
                interests = []
            interest_labels = ", ".join(INTEREST_LABELS[value] for value in interests if value in INTEREST_LABELS) if isinstance(interests, list) else ""
            sheet.append([
                record["created_at"], self._excel_text(record["name"]), self._excel_text(record["email"]),
                self._excel_text(record["phone"]), self._excel_text(record["postal_code"]),
                self._excel_text(record["city"]), self._excel_text(interest_labels), record["consent_at"], self._excel_text(record["consent_version"]),
            ])
        header_fill = PatternFill("solid", fgColor="102B46")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:I{max(1, sheet.max_row)}"
        for column, width in {"A": 25, "B": 28, "C": 36, "D": 22, "E": 16, "F": 26, "G": 34, "H": 25, "I": 20}.items():
            sheet.column_dimensions[column].width = width
        if records:
            table = Table(displayName="Besucherkontakte", ref=f"A1:I{sheet.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        workbook.properties.title = "AFS Besucherkontakte"
        workbook.properties.creator = "AFS contact intake"
        output.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
        output.parent.chmod(0o700)
        if output.is_symlink():
            raise ContactIntakeError("contact_export_unavailable")
        temporary = output.with_name(f".{output.stem}-{os.getpid()}.xlsx")
        try:
            workbook.save(temporary)
            temporary.chmod(0o600)
            os.replace(temporary, output)
            output.chmod(0o600)
        finally:
            temporary.unlink(missing_ok=True)

    @staticmethod
    def _git(repo: Path, *args: str) -> str:
        try:
            result = subprocess.run(
                ["git", *args], cwd=repo, capture_output=True, text=True, timeout=60,
                env={**os.environ, "GIT_TERMINAL_PROMPT": "0"}, check=True,
            )
        except (OSError, subprocess.SubprocessError):
            raise ContactIntakeError("contact_export_sync_failed") from None
        return result.stdout.strip()

    def export_and_push(self) -> None:
        if not self.enabled:
            raise ContactIntakeError("intake_unavailable")
        if not self.records():
            return
        config = self.config
        repo = Path(config["repositoryPath"]).expanduser()
        if not repo.is_absolute() or repo.is_symlink() or not (repo / ".git").is_dir():
            raise ContactIntakeError("contact_export_sync_failed")
        repo = repo.resolve()
        remote = self._git(repo, "remote", "get-url", "origin")
        parsed = urlsplit(remote)
        valid_remote = remote in {
            f"git@github.com:{REPOSITORY}.git",
            f"https://github.com/{REPOSITORY}.git",
        }
        if not valid_remote or parsed.username or parsed.password:
            raise ContactIntakeError("contact_export_sync_failed")
        branch = config.get("branch", "main")
        if branch != "main":
            raise ContactIntakeError("contact_export_sync_failed")
        export = repo / EXPORT_NAME
        self.export_workbook(export)
        status = self._git(repo, "status", "--porcelain", "--untracked-files=all")
        allowed = {EXPORT_NAME}
        if any(line.strip().split(maxsplit=1)[-1] not in allowed for line in status.splitlines() if line.strip()):
            raise ContactIntakeError("contact_export_sync_failed")
        self._git(repo, "add", "--", EXPORT_NAME)
        changed = self._git(repo, "diff", "--cached", "--name-only")
        if changed:
            if changed != EXPORT_NAME:
                raise ContactIntakeError("contact_export_sync_failed")
            self._git(repo, "-c", "user.name=AFS contact export", "-c", "user.email=afs-contact-export@localhost", "commit", "-m", "Update visitor contact export")
        self._git(repo, "push", "origin", "HEAD:main")
